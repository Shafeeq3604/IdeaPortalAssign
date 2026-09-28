"""Builds the SPEC §12.4 golden-set labelling workbooks from tests/evals/cases.ts (exported to JSON).

Outputs (into tests/evals/labelling/):
  - golden-set-annotator.xlsx      one copy per annotator: shuffled, neutral IDs, no groups, no draft labels
  - golden-set-adjudication.xlsx   for the person resolving disagreements: A vs B vs draft, per field
"""
import json
import random
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

SRC, OUT_DIR = sys.argv[1], sys.argv[2]
data = json.load(open(SRC, encoding="utf-8"))
vocab = data["vocab"]
cases = [c for c in data["cases"] if c["category"] != "adversarial"]
assert len(cases) == 40, len(cases)

# Deterministic shuffle so order does not reveal the SPEC group, and a rerun gives the same IDs.
rng = random.Random(20260924)
order = cases[:]
rng.shuffle(order)
for i, c in enumerate(order, start=1):
    c["label_id"] = f"L{i:02d}"

FONT = "Arial"
BASE = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True)
H2 = Font(name=FONT, size=11, bold=True)
WHITE_BOLD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")  # soft yellow = fill this in
GREY_FILL = PatternFill("solid", fgColor="F2F2F2")
DISAGREE_FILL = PatternFill("solid", fgColor="F8CBAD")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

VALUE_LABEL = {
    "BUSINESS_IMPACT": "Business impact", "PRODUCTIVITY": "Productivity", "COST_REDUCTION": "Cost reduction",
    "REVENUE": "Revenue", "EMPLOYEE_EXPERIENCE": "Employee experience", "CUSTOMER_IMPACT": "Customer impact",
    "OPERATIONAL": "Operational", "PROBLEM_SEVERITY": "Problem severity", "PROBLEM_FREQUENCY": "Problem frequency",
}
VALUE_HELP = {
    "BUSINESS_IMPACT": "Overall effect on how the organisation performs.",
    "PRODUCTIVITY": "Time or effort people get back.",
    "COST_REDUCTION": "Money the organisation stops spending.",
    "REVENUE": "Money the organisation newly earns.",
    "EMPLOYEE_EXPERIENCE": "How much better work feels for staff.",
    "CUSTOMER_IMPACT": "How much better things get for customers.",
    "OPERATIONAL": "Reliability, speed or quality of how work runs.",
    "PROBLEM_SEVERITY": "How bad the problem is when it happens.",
    "PROBLEM_FREQUENCY": "How often the problem happens.",
}
FEAS_HELP = {
    "HIGHLY_FEASIBLE": "Could be built with what the organisation plausibly has today.",
    "FEASIBLE_WITH_CONDITIONS": "Doable if named conditions are met (data access, a budget, a team).",
    "REQUIRES_INVESTIGATION": "Can't tell yet — a real unknown needs looking into first.",
    "NOT_CURRENTLY_FEASIBLE": "Blocked under the constraints the submission itself states.",
}
RISK_HELP = {
    "TECHNICAL": "It may not work, or be hard to build or run.",
    "SECURITY": "Could expose systems or data to attack or misuse.",
    "PRIVACY": "Touches personal data in a way that needs care.",
    "COMPLIANCE": "Regulatory, legal or policy obligations apply.",
    "FINANCIAL": "Cost overrun, or money lost if it goes wrong.",
    "OPERATIONAL": "Could disrupt how work runs day to day.",
    "ADOPTION": "People may not use it, or resist the change.",
    "DATA": "Needed data may be missing, poor or hard to get.",
    "VENDOR": "Depends on a supplier or third-party product.",
}
BAND_LABEL = {"NEGLIGIBLE": "Negligible", "LOW": "Low", "MODERATE": "Moderate", "HIGH": "High", "VERY_HIGH": "Very high"}

FIELD_ORDER = [
    ("title", "Title"), ("problemStatement", "Problem"), ("description", "Proposed idea"),
    ("expectedUsers", "Who would use it"), ("expectedOutcome", "Expected outcome"),
    ("existingProcess", "How it works today"), ("existingSolutions", "Existing solutions mentioned"),
    ("suggestedTechnology", "Suggested technology"), ("expectedBenefits", "Expected benefits"),
    ("estimatedCostNote", "Cost note"), ("references", "References"),
]

# Labels sheet columns — fixed, so the adjudication workbook and a later import can rely on them.
LABEL_COLS = (
    [("id", "ID"), ("title", "Idea title (for reference)"), ("useCases", "Use cases — one per line")]
    + [(f"value:{v}", f"Value — {VALUE_LABEL[v]}") for v in vocab["value"]]
    + [("feasibility", "Feasibility status")]
    + [(f"risk:{r}", f"Risk — {r.title()}") for r in vocab["risk"]]
    + [("notes", "Notes / doubts (optional)")]
)
COL = {key: i + 1 for i, (key, _) in enumerate(LABEL_COLS)}
LABEL_HEADER_ROW = 4
FIRST_DATA_ROW = LABEL_HEADER_ROW + 1


def style_header(cell):
    cell.font = WHITE_BOLD
    cell.fill = HEAD_FILL
    cell.alignment = CENTER
    cell.border = BOX


def write_lists(wb):
    """Dropdown sources on a hidden sheet (Excel list validations cap literal lists at 255 chars)."""
    ws = wb.create_sheet("Lists")
    ws["A1"], ws["B1"], ws["C1"] = "Band", "Feasibility", "Risk present"
    for i, b in enumerate(vocab["band"], start=2):
        ws.cell(i, 1, BAND_LABEL[b])
    for i, f in enumerate(vocab["feasibility"], start=2):
        ws.cell(i, 2, f)
    ws.cell(2, 3, "Yes")
    ws.sheet_state = "hidden"
    return (
        f"=Lists!$A$2:$A${1 + len(vocab['band'])}",
        f"=Lists!$B$2:$B${1 + len(vocab['feasibility'])}",
        "=Lists!$C$2:$C$2",
    )


def build_labels_sheet(wb, ws, band_ref, feas_ref, risk_ref, with_titles=True):
    ws["A1"] = "Annotator name:"
    ws["A1"].font = BOLD
    ws["B1"].fill = INPUT_FILL
    ws["B1"].border = BOX
    ws["D1"] = "Ideas with a feasibility status entered:"
    ws["D1"].font = BOLD
    last = FIRST_DATA_ROW + len(order) - 1
    fl = get_column_letter(COL["feasibility"])
    ws["H1"] = f'=COUNTA({fl}{FIRST_DATA_ROW}:{fl}{last})&" of {len(order)}"'
    ws["H1"].font = BOLD
    ws["A2"] = "Yellow cells are yours to fill in. Leave any cell blank if you cannot judge it from the text — blank is not wrong."
    ws["A2"].font = Font(name=FONT, size=9, italic=True)

    for key, header in LABEL_COLS:
        style_header(ws.cell(LABEL_HEADER_ROW, COL[key], header))
    for r, c in enumerate(order, start=FIRST_DATA_ROW):
        ws.cell(r, COL["id"], c["label_id"]).font = BOLD
        if with_titles:
            ws.cell(r, COL["title"], c["fields"].get("title") or "").font = BASE
        for key, _ in LABEL_COLS:
            cell = ws.cell(r, COL[key])
            cell.border = BOX
            cell.alignment = WRAP
            if key not in ("id", "title"):
                cell.fill = INPUT_FILL
                cell.font = BASE
    rng_for = lambda key: f"{get_column_letter(COL[key])}{FIRST_DATA_ROW}:{get_column_letter(COL[key])}{last}"
    dv_band = DataValidation(type="list", formula1=band_ref, allow_blank=True,
                             error="Pick a band from the list, or leave it blank.", errorTitle="Not a band")
    dv_feas = DataValidation(type="list", formula1=feas_ref, allow_blank=True,
                             error="Pick a status from the list, or leave it blank.", errorTitle="Not a status")
    dv_risk = DataValidation(type="list", formula1=risk_ref, allow_blank=True,
                             error='Choose "Yes", or leave it blank.', errorTitle="Yes or blank")
    for dv in (dv_band, dv_feas, dv_risk):
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
    for v in vocab["value"]:
        dv_band.add(rng_for(f"value:{v}"))
    dv_feas.add(rng_for("feasibility"))
    for rk in vocab["risk"]:
        dv_risk.add(rng_for(f"risk:{rk}"))

    widths = {"id": 7, "title": 34, "useCases": 42, "feasibility": 26, "notes": 34}
    for key, _ in LABEL_COLS:
        w = widths.get(key, 13 if key.startswith("value:") else 11)
        ws.column_dimensions[get_column_letter(COL[key])].width = w
    ws.row_dimensions[LABEL_HEADER_ROW].height = 42
    for r in range(FIRST_DATA_ROW, last + 1):
        ws.row_dimensions[r].height = 60
    ws.freeze_panes = ws.cell(FIRST_DATA_ROW, COL["useCases"])
    return last


def annotator_workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 26
    rows = [
        ("Golden-set labelling", TITLE),
        ("For the Idea Platform's AI evaluation (SPEC §12.4). You are one of two independent annotators.", BASE),
        ("", None),
        ("What to do", H2),
        ("1. Read each idea on the Ideas tab. Judge it only from what is written there.", BASE),
        ("2. Fill in its row on the Labels tab: use cases, a band per value dimension, a feasibility status, and which risks apply.", BASE),
        ("3. Leave anything blank that you cannot confidently judge from the text. A blank is excluded from scoring, never counted as wrong.", BASE),
        ("4. Work alone. Do not compare answers with the other annotator — the whole point is two independent reads.", BASE),
        ("5. Put your name in cell B1 of the Labels tab, save, and send the file back.", BASE),
        ("", None),
        ("Legend", H2),
    ]
    r = 1
    for text, font in rows:
        ws.cell(r, 1, text).font = font or BASE
        r += 1
    ws.cell(r, 1, "Yellow cell").fill = INPUT_FILL
    ws.cell(r, 1).border = BOX
    ws.cell(r, 2, "Yours to fill in (Labels tab). Everything else is read-only reference.").font = BASE
    r += 2

    def table(title, header, items):
        nonlocal r
        ws.cell(r, 1, title).font = H2
        r += 1
        for i, h in enumerate(header, start=1):
            style_header(ws.cell(r, i, h))
        r += 1
        for vals in items:
            for i, v in enumerate(vals, start=1):
                c = ws.cell(r, i, v)
                c.font = BASE
                c.alignment = WRAP
                c.border = BOX
            r += 1
        r += 1

    table("Use cases", ["Guidance", "Detail"], [
        ("What counts", "A concrete way people would use the idea — a situation, not a benefit. \"Remind an approver after 24 hours of inactivity\" is a use case; \"saves time\" is not."),
        ("How to write them", "One per line in the cell (Alt+Enter for a new line). Short phrases are fine — wording does not need to match anything exactly."),
    ])
    table("Value bands — one per dimension", ["Dimension", "What it means", "Bands"],
          [(VALUE_LABEL[v], VALUE_HELP[v], "Negligible · Low · Moderate · High · Very high") for v in vocab["value"]])
    table("Feasibility status — pick one", ["Status", "When to use it"], [(f, FEAS_HELP[f]) for f in vocab["feasibility"]])
    table("Risks — mark \"Yes\" for each one a careful reviewer should raise", ["Risk", "What it means"],
          [(rk.title(), RISK_HELP[rk]) for rk in vocab["risk"]])

    ws.cell(r, 1, "Example row (a made-up idea — not in the set)").font = H2
    r += 1
    example = {
        "id": "EX", "title": "Auto-release unused meeting rooms",
        "useCases": "Release a booked room if nobody checks in within 10 minutes\nShow freed rooms to people searching now",
        "value:PRODUCTIVITY": "Moderate", "value:PROBLEM_FREQUENCY": "High", "value:EMPLOYEE_EXPERIENCE": "Moderate",
        "feasibility": "FEASIBLE_WITH_CONDITIONS", "risk:ADOPTION": "Yes", "risk:TECHNICAL": "Yes",
        "notes": "Conditional on room sensors or a check-in step existing.",
    }
    shown = [("id", "ID"), ("title", "Idea title"), ("useCases", "Use cases"), ("value:PRODUCTIVITY", "Value — Productivity"),
             ("value:PROBLEM_FREQUENCY", "Value — Problem frequency"), ("feasibility", "Feasibility status"),
             ("risk:ADOPTION", "Risk — Adoption"), ("notes", "Notes")]
    for i, (_, h) in enumerate(shown, start=1):
        style_header(ws.cell(r, i, h))
    r += 1
    for i, (k, _) in enumerate(shown, start=1):
        c = ws.cell(r, i, example.get(k, ""))
        c.font = BASE
        c.alignment = WRAP
        c.border = BOX
        c.fill = GREY_FILL
    ws.row_dimensions[r].height = 48
    ws.cell(r + 1, 1, "Other value dimensions left blank in the example = not judgeable from that text. That is normal.").font = Font(name=FONT, size=9, italic=True)
    for col in "DEFGH":
        ws.column_dimensions[col].width = 16

    ideas = wb.create_sheet("Ideas")
    ideas["A1"] = "The 40 ideas, in a fixed shuffled order. IDs match the Labels tab."
    ideas["A1"].font = BOLD
    headers = [("ID", 7)] + [(label, 40 if key in ("problemStatement", "description") else 28) for key, label in FIELD_ORDER]
    for i, (h, w) in enumerate(headers, start=1):
        style_header(ideas.cell(3, i, h))
        ideas.column_dimensions[get_column_letter(i)].width = w
    for rr, c in enumerate(order, start=4):
        ideas.cell(rr, 1, c["label_id"]).font = BOLD
        for i, (key, _) in enumerate(FIELD_ORDER, start=2):
            cell = ideas.cell(rr, i, c["fields"].get(key) or "—")
            cell.font = BASE
            cell.alignment = WRAP
            cell.border = BOX
        ideas.row_dimensions[rr].height = 150
    ideas.freeze_panes = "C4"

    band_ref, feas_ref, risk_ref = write_lists(wb)
    labels = wb.create_sheet("Labels", 2)
    build_labels_sheet(wb, labels, band_ref, feas_ref, risk_ref)
    wb.active = 0
    wb.save(path)


def adjudication_workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    ws.column_dimensions["A"].width = 110
    lines = [
        ("Golden-set adjudication", TITLE),
        ("For the person resolving disagreements between the two annotators (SPEC §12.4: \"two annotators with disagreements resolved\").", BASE),
        ("", None),
        ("1. Open each annotator's returned file, select their whole Labels tab (Ctrl+A), copy, and paste into cell A1 of 'Annotator A' / 'Annotator B' here.", BASE),
        ("2. The Compare tab lines up A, B and the original draft label for every field. Rows where A and B differ are shaded.", BASE),
        ("3. For each shaded row, decide and type the final label in the yellow 'Agreed' column. Where A and B agree, 'Agreed' may be left blank — the agreed value is used.", BASE),
        ("4. Use cases are free text, so they are never auto-compared: read both and write the agreed list.", BASE),
        ("5. Send this file back. The agreed labels replace the draft groundTruth blocks in tests/evals/cases.ts.", BASE),
        ("", None),
        ("The 'Draft' column is the single-author first pass already in the code. It is shown only here, never to the annotators, so it could not bias their reads.", Font(name=FONT, size=9, italic=True)),
        ("The Key tab maps each L-ID to its case in the code and its SPEC group. Annotators never saw the groups.", Font(name=FONT, size=9, italic=True)),
    ]
    for i, (t, f) in enumerate(lines, start=1):
        ws.cell(i, 1, t).font = f or BASE
        ws.cell(i, 1).alignment = Alignment(wrap_text=True)

    band_ref, feas_ref, risk_ref = write_lists(wb)
    for name in ("Annotator A", "Annotator B"):
        s = wb.create_sheet(name)
        build_labels_sheet(wb, s, band_ref, feas_ref, risk_ref)

    key = wb.create_sheet("Key")
    for i, h in enumerate(["ID", "Case name in tests/evals/cases.ts", "SPEC group", "Title"], start=1):
        style_header(key.cell(1, i, h))
    for rr, c in enumerate(order, start=2):
        for i, v in enumerate([c["label_id"], c["name"], c["category"], c["fields"].get("title") or ""], start=1):
            key.cell(rr, i, v).font = BASE
    for col, w in zip("ABCD", (7, 60, 16, 50)):
        key.column_dimensions[col].width = w
    key.freeze_panes = "A2"

    cmp_ = wb.create_sheet("Compare", 1)
    heads = [("ID", 7), ("SPEC group", 14), ("Field", 26), ("Draft (code today)", 24), ("Annotator A", 24),
             ("Annotator B", 24), ("A vs B", 11), ("Agreed", 26), ("Final label", 24)]
    for i, (h, w) in enumerate(heads, start=1):
        style_header(cmp_.cell(1, i, h))
        cmp_.column_dimensions[get_column_letter(i)].width = w
    fields = [("useCases", "Use cases")] + [(f"value:{v}", f"Value — {VALUE_LABEL[v]}") for v in vocab["value"]] \
        + [("feasibility", "Feasibility status")] + [(f"risk:{rk}", f"Risk — {rk.title()}") for rk in vocab["risk"]]
    last_label_row = FIRST_DATA_ROW + len(order) - 1
    row = 2
    for c in order:
        gt = c["groundTruth"] or {}
        for fkey, flabel in fields:
            if fkey == "useCases":
                draft = "\n".join(gt.get("useCases", []))
            elif fkey.startswith("value:"):
                b = (gt.get("valueBands") or {}).get(fkey.split(":")[1])
                draft = BAND_LABEL[b] if b else ""
            elif fkey == "feasibility":
                draft = gt.get("feasibilityStatus", "")
            else:
                draft = "Yes" if fkey.split(":")[1] in (gt.get("riskCategories") or []) else ""
            colL = get_column_letter(COL[fkey])
            idref = f"$A{row}"

            def pick(sheet):
                rng_ids = f"'{sheet}'!$A${FIRST_DATA_ROW}:$A${last_label_row}"
                rng_val = f"'{sheet}'!${colL}${FIRST_DATA_ROW}:${colL}${last_label_row}"
                return f'=IFERROR(INDEX({rng_val},MATCH({idref},{rng_ids},0))&"","")'

            vals = [c["label_id"], c["category"], flabel, draft, pick("Annotator A"), pick("Annotator B")]
            if fkey == "useCases":
                vals.append("read both")
            else:
                vals.append(f'=IF(AND(E{row}="",F{row}=""),"",IF(E{row}=F{row},"agree","DIFFER"))')
            vals.append(None)
            if fkey == "useCases":
                vals.append(f'=IF(H{row}<>"",H{row},"")')
            else:
                vals.append(f'=IF(H{row}<>"",H{row},IF(G{row}="agree",E{row},""))')
            for i, v in enumerate(vals, start=1):
                cell = cmp_.cell(row, i, v)
                cell.font = BASE
                cell.alignment = WRAP
                cell.border = BOX
            cmp_.cell(row, 8).fill = INPUT_FILL
            row += 1
    last_cmp = row - 1
    cmp_.conditional_formatting.add(f"A2:I{last_cmp}", FormulaRule(formula=[f'$G2="DIFFER"'], fill=DISAGREE_FILL))
    dv_feas = DataValidation(type="list", formula1=feas_ref, allow_blank=True)
    cmp_.freeze_panes = "D2"
    cmp_.auto_filter.ref = f"A1:I{last_cmp}"

    summary_row = 1
    cmp_.cell(summary_row, 11, "Rows where A and B differ:").font = BOLD
    cmp_.cell(summary_row, 12, f'=COUNTIF(G2:G{last_cmp},"DIFFER")').font = BOLD
    cmp_.cell(summary_row + 1, 11, "Differing rows still without an agreed label:").font = BOLD
    cmp_.cell(summary_row + 1, 12, f'=COUNTIFS(G2:G{last_cmp},"DIFFER",H2:H{last_cmp},"")').font = BOLD
    cmp_.column_dimensions["K"].width = 40
    wb.active = 0
    wb.save(path)


annotator_workbook(f"{OUT_DIR}/golden-set-annotator.xlsx")
adjudication_workbook(f"{OUT_DIR}/golden-set-adjudication.xlsx")
print("ok", len(order), "ideas")
